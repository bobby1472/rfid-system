"""สร้างไฟล์ Excel ของประวัติการเข้าออก

แยกออกมาจาก router เพื่อให้ตรรกะการจัดหน้าตาไฟล์ไม่ปนกับเรื่อง HTTP
"""
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import AccessLog

# (หัวคอลัมน์, ความกว้าง)
COLUMNS = [
    ("ลำดับ", 8),
    ("วันที่ - เวลา", 21),
    ("เครื่องอ่าน", 13),
    ("UID", 16),
    ("บาร์โค้ด", 20),
    ("ชื่อผู้ถือบัตร", 34),
    ("ชนิดบัตร", 14),
    ("ผล", 10),
    ("เหตุผล", 22),
]

REASON_TH = {
    "ok": "ผ่าน",
    "unknown_card": "บัตรไม่อยู่ในระบบ",
    "card_disabled": "บัตรถูกปิดใช้งาน",
    "reader_disabled": "เครื่องถูกปิดใช้งาน",
}

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_DENY_FILL = PatternFill("solid", fgColor="FCE4E4")


def build_workbook(rows: list[AccessLog], tz_name: str) -> BytesIO:
    try:
        tz = ZoneInfo(tz_name)
    except Exception:
        tz = ZoneInfo("UTC")

    wb = Workbook()
    ws = wb.active
    ws.title = "Access Logs"

    for col, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    for i, log in enumerate(rows, start=1):
        # เวลาในฐานข้อมูลเก็บเป็น UTC ต้องแปลงเป็นเวลาท้องถิ่นก่อนเขียนลงไฟล์
        # แล้วตัด tzinfo ทิ้ง เพราะ Excel ไม่รองรับ datetime ที่มีเขตเวลาติดมา
        local_dt = log.scanned_at.astimezone(tz).replace(tzinfo=None) if log.scanned_at else None

        values = [
            i,
            local_dt,
            log.reader_code or "-",
            log.uid,
            log.barcode or "-",
            log.holder_name or "ไม่รู้จัก",
            log.card_type or "-",
            "ผ่าน" if log.allowed else "ปฏิเสธ",
            REASON_TH.get(log.reason, log.reason),
        ]

        r = i + 1
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            if col == 2:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            if col in (1, 3, 7, 8):
                cell.alignment = Alignment(horizontal="center")
            # ระบายแถวที่ถูกปฏิเสธให้เห็นชัดตอนเปิดดู
            if not log.allowed:
                cell.fill = _DENY_FILL

    # ตรึงหัวตารางไว้ และเปิดตัวกรองให้พร้อมใช้
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def filename(tz_name: str) -> str:
    try:
        tz = ZoneInfo(tz_name)
    except Exception:
        tz = ZoneInfo("UTC")
    return f"access_logs_{datetime.now(tz):%Y%m%d_%H%M}.xlsx"
